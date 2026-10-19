"""卖家精灵原始关键词Excel清洗；仅接受可核实的美国站ASIN快照。"""
import json
import re
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


class ExcelValidationError(ValueError):
    pass


# 使用双语名称映射，不依赖列序号；更新时间由其前面的排名列确定含义。
ALIASES = {
    'keyword': ('关键词', 'Keyword Phrase'),
    'translation': ('关键词翻译', 'Keyword Translation'),
    'traffic_share': ('流量占比', 'Click Share'),
    'weekly_impressions': ('预估周曝光量', 'Estimated Weekly Impressions'),
    'keyword_type': ('关键词类型', 'Keywords type'),
    'conversion_label': ('转化效果', 'Conversion'),
    'traffic_type': ('流量词类型', 'Keyword Distribution'),
    'organic_share': ('自然流量占比', 'Organic Share'),
    'sponsored_share': ('广告流量占比', 'Sponsored Share'),
    'organic_rank': ('自然排名', 'Organic Rank'),
    'organic_page': ('自然排名页码', 'Organic page'),
    'sponsored_rank': ('广告排名', 'Sponsored Rank'),
    'sponsored_page': ('广告排名页码', 'Sponsored page'),
    'aba_week_rank': ('ABA周排名', 'Weekly ABA Rank'),
    'monthly_searches': ('月搜索量', 'Monthly Searches'),
    'search_growth_mom': ('搜索量环比增长率', 'MoM Search Growth Rate'),
    'search_growth_yoy': ('搜索量同比增长率', 'YoY Search Growth Rate'),
    'spr': ('SPR',),
    'title_density': ('标题密度', 'Title density'),
    'monthly_purchases': ('购买量', '月购买量', 'Monthly Purchase'),
    'purchase_rate': ('购买率', 'Purchase Rate'),
    'impressions': ('展示量', 'Impressions'),
    'clicks': ('点击量', 'Clicks'),
    'products': ('商品数', 'Products'),
    'demand_supply_ratio': ('需供比', 'DSR'),
    'sponsored_asins': ('近7天广告竞品数', 'Sponsored ASINs'),
    'total_click_share': ('点击总占比', 'Total Click Share'),
    'total_conversion_share': ('转化总占比', 'Total Conversion Share'),
    'ppc': ('PPC价格', 'PPC'),
    'suggested_bid': ('建议竞价范围', 'Suggested bid'),
    'top_asins': ('前十ASIN', 'TOP 10 Products'),
}
INTEGER_FIELDS = {
    'weekly_impressions', 'organic_rank', 'sponsored_rank', 'aba_week_rank',
    'monthly_searches', 'spr', 'title_density', 'monthly_purchases', 'impressions',
    'clicks', 'products', 'sponsored_asins',
}
DECIMAL_FIELDS = {
    'traffic_share', 'organic_share', 'sponsored_share', 'search_growth_mom',
    'search_growth_yoy', 'purchase_rate', 'demand_supply_ratio', 'total_click_share',
    'total_conversion_share', 'ppc',
}
TEXT_FIELDS = {'translation', 'keyword_type', 'conversion_label', 'traffic_type',
               'organic_page', 'sponsored_page', 'organic_updated_text', 'sponsored_updated_text'}


def normalized(value):
    return re.sub(r'\s+', '', str(value or '')).casefold()


HEADER_MAP = {normalized(alias): field for field, names in ALIASES.items() for alias in names}
MISSING = {'', '-', '--', 'none', 'null', 'n/a', 'na', '前3页无排名', '暂无数据'}


def text_value(value):
    if value is None or str(value).strip().casefold() in MISSING:
        return None
    return str(value).strip()


def number_value(value, integer=False, signed=False):
    raw = text_value(value)
    if raw is None:
        return None
    percent = raw.endswith('%')
    raw = re.sub(r'[$,\s]', '', raw.rstrip('%'))
    try:
        number = Decimal(raw)
    except InvalidOperation as exc:
        raise ExcelValidationError('无法识别数字字段') from exc
    if not number.is_finite() or (not signed and number < 0):
        raise ExcelValidationError('数字字段超出允许范围')
    if percent:
        number /= 100
    if integer:
        if number != number.to_integral_value() or number > 9223372036854775807:
            raise ExcelValidationError('计数或排名不是有效整数')
        return int(number)
    if abs(number) >= Decimal('1000000000000'):
        raise ExcelValidationError('数值超出数据库精度')
    return number.quantize(Decimal('0.00000001'))


def parse_excel(path, asin, marketplace='US'):
    """提取关键词主表；空值、真零、原始双时区排名时间分别保留。"""
    path = Path(path)
    if marketplace != 'US' or not re.fullmatch(r'[A-Z0-9]{10}', asin):
        raise ExcelValidationError('仅支持合法的美国站ASIN')
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if (len(members) > 5000 or sum(m.file_size for m in members) > 150 * 1024 * 1024
                    or 'xl/workbook.xml' not in archive.namelist()):
                raise ExcelValidationError('文件结构或展开体积超出限制')
    except zipfile.BadZipFile as exc:
        raise ExcelValidationError('不是有效的xlsx工作簿') from exc
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        candidates = []
        for sheet in workbook:
            if sheet.max_row > 50002 or sheet.max_column > 100:
                raise ExcelValidationError('工作表超过行列限制')
            headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            if any(HEADER_MAP.get(normalized(h)) == 'keyword' for h in headers):
                candidates.append((sheet, headers))
        if len(candidates) != 1:
            raise ExcelValidationError('必须有且仅有一个关键词主表')
        sheet, headers = candidates[0]
        identity = re.search(r'(US|美国)[-_]([A-Z0-9]{10})', sheet.title, re.I)
        if not identity or identity.group(2).upper() != asin:
            raise ExcelValidationError('工作表中的站点或ASIN与上传任务不一致')
        mapping = {}
        raw_names = []
        occurrences = {}
        context = None
        for index, header in enumerate(headers):
            name = str(header or f'column_{index + 1}').strip()
            occurrences[name] = occurrences.get(name, 0) + 1
            raw_names.append(name if occurrences[name] == 1 else f'{name}_{occurrences[name]}')
            field = HEADER_MAP.get(normalized(header))
            if field in ('organic_rank', 'sponsored_rank'):
                context = field.split('_')[0]
            if normalized(header) in ('updated', '更新时间') and context:
                field = f'{context}_updated_text'
            if field:
                if field in mapping.values():
                    raise ExcelValidationError('主表存在重复业务字段')
                mapping[index] = field
        required = {'keyword', 'organic_rank', 'sponsored_rank', 'monthly_searches', 'traffic_share'}
        if not required.issubset(mapping.values()):
            raise ExcelValidationError('关键词主表缺少必要字段，可能导出模板已改变')
        rows, seen = [], set()
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            source = {field: values[index] for index, field in mapping.items()}
            keyword = text_value(source.get('keyword'))
            if not keyword:
                if any(text_value(v) for v in values):
                    raise ExcelValidationError(f'第{row_number}行有数据但缺少关键词')
                continue
            if len(keyword) > 500:
                raise ExcelValidationError(f'第{row_number}行关键词过长')
            key = re.sub(r'\s+', ' ', keyword).casefold()
            if key in seen:
                raise ExcelValidationError('关键词重复，拒绝静默覆盖')
            seen.add(key)
            row = {'keyword': keyword, 'keyword_normalized': key, 'source_row': row_number}
            for field in INTEGER_FIELDS:
                row[field] = number_value(source.get(field), integer=True)
            for field in DECIMAL_FIELDS:
                row[field] = number_value(source.get(field), signed=field.startswith('search_growth_'))
            for field in TEXT_FIELDS:
                row[field] = text_value(source.get(field))
                if row[field] and len(row[field]) > 1000:
                    raise ExcelValidationError(f'第{row_number}行文本字段过长')
            bid = text_value(source.get('suggested_bid'))
            row['suggested_bid_min'] = row['suggested_bid_max'] = None
            if bid:
                match = re.fullmatch(r'\$?([\d,.]+)\s*[-–~]\s*\$?([\d,.]+)', bid)
                if not match:
                    raise ExcelValidationError('建议竞价范围格式未知')
                row['suggested_bid_min'], row['suggested_bid_max'] = map(number_value, match.groups())
                if row['suggested_bid_min'] > row['suggested_bid_max']:
                    raise ExcelValidationError('建议竞价范围上下限错误')
            top_asins = re.findall(r'\b[A-Z0-9]{10}\b', str(source.get('top_asins') or '').upper())
            row['top_asins_json'] = json.dumps(top_asins)
            row['raw_row_json'] = json.dumps(dict(zip(raw_names, values)), ensure_ascii=False, default=str)
            rows.append(row)
        count = re.search(r'\((\d+)\)', sheet.title)
        if count and int(count.group(1)) != len(rows):
            raise ExcelValidationError('导出数量与主表记录数不一致')
        if not rows:
            raise ExcelValidationError('空关键词表，不能标记清洗成功')
        metadata = {'sheet_name': sheet.title, 'headers': list(headers), 'row_count': len(rows), 'word_frequency': []}
        for other in workbook:
            if other.title in ('Unique Words', '高频词', '词频统计'):
                metadata['word_frequency'] = [list(v) for v in other.iter_rows(min_row=2, values_only=True)
                                               if any(x is not None for x in v)]
        return rows, metadata
    finally:
        workbook.close()
