"""查竞品普通导出清洗：逐行校验批量ASIN集合，所有原始字段完整留存。"""
import json
import re
import zipfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from services.sellersprite_excel_common import ExcelValidationError, number_value, text_value

# 普通导出66列中的常用查询指标；其余列仍保留在详情JSON和原始行JSON。
FIELDS = {
 'asin':('ASIN',), 'parent_asin':('Parent','父ASIN','父体ASIN'),
 'title':('Product Title','产品标题','商品标题','标题'), 'brand':('Brand','品牌'),
 'product_url':('URL','产品链接','商品链接'), 'image_url':('Product Image','产品图片','主图'),
 'category_path':('Category Path','类目路径','所属类目'), 'category':('Category','大类目'),
 'bsr':('Category BSR','大类BSR'), 'bsr_change':('Category BSR Increase','大类BSR增长数'),
 'bsr_growth':('Category BSR Growth','大类BSR增长率'),
 'subcategory':('Sub-Category','小类目','子类目'), 'subcategory_bsr':('Sub-Category BSR','小类BSR'),
 'monthly_sales':('Sales','月销量','销量'), 'sales_growth_mom':('MoM Sales Growth Rate','销量环比增长率'),
 'sales_growth_yoy':('YoY Sales Growth Rate','销量同比增长率'),
 'monthly_revenue':('Monthly Revenue($)','月销售额($)','月销售额($)','月销售额'),
 'variation_sales':('Variation Sold','子体销量'), 'variation_revenue':('Variation Revenue($)','子体销售额($)'),
 'variation_count':('Variations','变体数'), 'price':('Price($)','价格($)','价格'),
 'prime_price':('Prime Price($)','Prime价格($)'), 'ratings':('Ratings','评分数'),
 'rating_increase':('Monthly Rating Increase','月新增评分数'), 'rating':('Rating','评分'),
 'reviews_rate':('Reviews Rate','留评率'), 'fba_fee':('FBA($)','FBA费用($)','FBA($)'),
 'gross_margin':('Gross Margin','毛利率'), 'date_available':('Date Available','上架时间'),
 'lqs':('LQS',), 'sellers':('Sellers','卖家数'), 'seller_type':('Seller Type','配送方式'),
 'buybox_seller':('BuyBox Seller','BuyBox卖家'), 'buybox_location':('BuyBox Location','BuyBox卖家所属地'),
}
INT_FIELDS={'bsr','bsr_change','subcategory_bsr','monthly_sales','variation_sales','variation_count','ratings','rating_increase','sellers'}
DEC_FIELDS={'bsr_growth','sales_growth_mom','sales_growth_yoy','monthly_revenue','variation_revenue','price','prime_price','rating','reviews_rate','fba_fee','gross_margin','lqs'}
SIGNED={'bsr_change','bsr_growth','sales_growth_mom','sales_growth_yoy','rating_increase','gross_margin'}
FLAG_HEADERS={"Best Seller","Amazon's Choice",'New Release','EBC Available','Video Available','Sponsored Product','Brand Story','Sponsored Brands','LTD'}


def header(value):
    return re.sub(r'\s+','',str(value or '')).casefold().replace('（','(').replace('）',')').replace('＄','$')


def parse_report(path, selected_asin):
    """兼容原单商品调用，不接受多商品冒充单条导出。"""
    rows,metadata=parse_reports(path,[selected_asin])
    metadata['selected_asin']=selected_asin
    return rows[selected_asin],metadata


def parse_reports(path, selected_asins):
    """完整验证集合与行数后再按ASIN拆分，拒绝漏项、多项和重复行。"""
    if not isinstance(selected_asins,list) or not 1<=len(selected_asins)<=60 or len(set(selected_asins))!=len(selected_asins) or any(not re.fullmatch('[A-Z0-9]{10}',str(a or '')) for a in selected_asins):
        raise ExcelValidationError('必须提供合法且不重复的勾选ASIN集合')
    try:
        with zipfile.ZipFile(path) as z:
            if 'xl/workbook.xml' not in z.namelist() or len(z.infolist())>5000 or sum(i.file_size for i in z.infolist())>150*1024*1024:
                raise ExcelValidationError('报告格式或展开体积超限')
    except zipfile.BadZipFile as exc:
        raise ExcelValidationError('报告不是xlsx') from exc
    book=load_workbook(Path(path),read_only=True,data_only=True,keep_links=False)
    try:
        candidates=[]
        for sheet in book:
            if sheet.max_row>10000 or sheet.max_column>120:
                raise ExcelValidationError('报告行列超限')
            heads=next(sheet.values,())
            if 'asin' in [header(h) for h in heads]: candidates.append((sheet,heads))
        if len(candidates)!=1: raise ExcelValidationError('找不到唯一商品主表')
        sheet,heads=candidates[0]
        if any(not h for h in heads) or len(set(map(header,heads)))!=len(heads):
            raise ExcelValidationError('报告存在空白或重复列名')
        lookup={header(alias):key for key,names in FIELDS.items() for alias in names}
        mapping={index:lookup[header(h)] for index,h in enumerate(heads) if header(h) in lookup}
        if not {'asin','title','monthly_sales','price','bsr'}.issubset(mapping.values()):
            raise ExcelValidationError('不是查竞品普通导出模板或缺少关键字段')
        records=[values for values in sheet.iter_rows(min_row=2,values_only=True) if any(text_value(x) is not None for x in values)]
        if len(records)!=len(selected_asins):raise ExcelValidationError('报告行数与实际勾选数量不一致')
        if 'US' not in sheet.title.upper():raise ExcelValidationError('报告不是美国站')
        result={}
        for values in records:
            source={field:values[index] for index,field in mapping.items()}
            asin=str(source.get('asin') or '').strip().upper()
            if asin not in selected_asins or asin in result:raise ExcelValidationError('报告ASIN集合不一致或出现重复商品')
            result[asin]=clean_product(source,heads,values)
        if set(result)!=set(selected_asins):raise ExcelValidationError('报告遗漏已勾选ASIN')
        return result,{'sheet_name':sheet.title,'headers':list(heads),'row_count':len(result),'selected_asins':selected_asins,'currency':'USD','period':'last30days'}
    finally: book.close()


def clean_product(source,heads,values):
    """每条商品沿用原数值/空值/单位合同，不能取报告首行套给全部任务。"""
    url=text_value(source.get('product_url'))
    if url and not re.match(r'https://(?:www\.)?amazon\.com/',url):raise ExcelValidationError('商品链接不是美国站')
    row={}
    for field in FIELDS:
        value=source.get(field)
        if field in INT_FIELDS:row[field]=number_value(value,integer=True,signed=field in SIGNED)
        elif field in DEC_FIELDS:row[field]=number_value(value,signed=field in SIGNED)
        elif field=='date_available':
            value=text_value(value)
            if value:
                if isinstance(source[field],(datetime,date)):value=source[field].date().isoformat() if isinstance(source[field],datetime) else source[field].isoformat()
                else:
                    try:value=date.fromisoformat(value[:10]).isoformat()
                    except ValueError as exc:raise ExcelValidationError('上架日期格式无法识别') from exc
            row[field]=value
        else:row[field]=text_value(value)
    row['asin']=str(row['asin']).strip().upper()
    if row['rating'] is not None and row['rating']>5:raise ExcelValidationError('评分超出0～5')
    for key,value in row.items():
        if isinstance(value,str) and len(value)>(4000 if key in ('title','product_url','image_url','category_path') else 1000):raise ExcelValidationError('商品字段过长')
    raw=dict(zip(map(str,heads),values));details={}
    for name,value in raw.items():
        cleaned=text_value(value)
        if name in FLAG_HEADERS and cleaned is not None:
            flag=cleaned.casefold()
            if flag in ('y','yes','true','1','是'):cleaned=True
            elif flag in ('n','no','false','0','否'):cleaned=False
        details[name]=cleaned
    row['details_json']=json.dumps(details,ensure_ascii=False,default=str)
    row['raw_row_json']=json.dumps(raw,ensure_ascii=False,default=str)
    return row
